"""Local document text/layout extraction from uploaded bytes (no external LLM)."""

from __future__ import annotations

import io

from dar.providers.ports import BBox, OcrDocumentResult, OcrPageResult, OcrSpan

MAX_PAGES_DEFAULT = 50


def _lines_to_page(
    *,
    page_number: int,
    width: float,
    height: float,
    lines: list[tuple[str, float, float, float, float]],
    source: str,
    confidence: float,
    language: str = "ru",
) -> OcrPageResult:
    line_spans: list[OcrSpan] = []
    texts: list[str] = []
    for text, x, y, w, h in lines:
        t = text.strip()
        if not t:
            continue
        texts.append(t)
        line_spans.append(OcrSpan(text=t, bbox=BBox(x=x, y=y, w=w, h=h), confidence=confidence, language=language))
    full = "\n".join(texts)
    return OcrPageResult(
        page_number=page_number,
        text=full,
        confidence=confidence,
        width=width,
        height=height,
        rotation=0,
        language=language,
        words=tuple(line_spans),
        lines=tuple(line_spans),
        blocks=(),
        source=source,
        error_code=None if full else "empty_page",
    )


def _error_page(code: str) -> OcrPageResult:
    return OcrPageResult(
        page_number=1,
        text="",
        confidence=0.0,
        width=595.0,
        height=842.0,
        rotation=0,
        language="ru",
        words=(),
        lines=(),
        blocks=(),
        source="unknown",
        error_code=code,
    )


def _extract_pdf_pymupdf(data: bytes, max_pages: int) -> list[OcrPageResult]:
    import fitz  # pymupdf

    doc = fitz.open(stream=data, filetype="pdf")
    pages: list[OcrPageResult] = []
    for i in range(min(len(doc), max_pages)):
        page = doc.load_page(i)
        rect = page.rect
        blocks = page.get_text("dict").get("blocks", [])
        lines: list[tuple[str, float, float, float, float]] = []
        for block in blocks:
            if block.get("type") != 0:
                continue
            for line in block.get("lines", []):
                spans = line.get("spans", [])
                if not spans:
                    continue
                text = "".join(s.get("text", "") for s in spans).strip()
                if not text:
                    continue
                bbox = line.get("bbox", spans[0].get("bbox", (0, 0, 0, 0)))
                x0, y0, x1, y1 = bbox
                lines.append((text, float(x0), float(y0), max(1.0, float(x1 - x0)), max(1.0, float(y1 - y0))))
        has_native = bool(lines)
        if not has_native:
            raw = page.get_text("text") or ""
            y = 72.0
            for ln in raw.splitlines():
                t = ln.strip()
                if t:
                    lines.append((t, 72.0, y, 450.0, 14.0))
                    y += 16.0
            has_native = bool(lines)
        conf = 0.95 if has_native else 0.35
        source = "native_pdf" if has_native else "scan"
        pages.append(
            _lines_to_page(
                page_number=i + 1,
                width=float(rect.width),
                height=float(rect.height),
                lines=lines,
                source=source,
                confidence=conf,
            ),
        )
    doc.close()
    return pages


def _extract_pdf_pypdf(data: bytes, max_pages: int) -> list[OcrPageResult]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages: list[OcrPageResult] = []
    for i, page in enumerate(reader.pages[:max_pages]):
        text = page.extract_text() or ""
        lines_raw = [ln for ln in text.splitlines() if ln.strip()]
        y = 72.0
        lines: list[tuple[str, float, float, float, float]] = []
        for ln in lines_raw:
            lines.append((ln, 72.0, y, 450.0, 14.0))
            y += 16.0
        pages.append(
            _lines_to_page(
                page_number=i + 1,
                width=595.0,
                height=842.0,
                lines=lines,
                source="native_pdf" if lines else "scan",
                confidence=0.88 if lines else 0.3,
            ),
        )
    return pages


def _extract_docx(data: bytes, max_pages: int) -> list[OcrPageResult]:
    try:
        from docx import Document
    except ImportError:
        return [_error_page("missing_docx")]

    doc = Document(io.BytesIO(data))
    lines: list[tuple[str, float, float, float, float]] = []
    y = 72.0
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            lines.append((t, 72.0, y, 450.0, 14.0))
            y += 16.0
    for table in doc.tables:
        for row in table.rows:
            cells = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
            if cells:
                lines.append((cells, 72.0, y, 450.0, 14.0))
                y += 16.0
    return [
        _lines_to_page(
            page_number=1,
            width=595.0,
            height=max(842.0, y + 40),
            lines=lines[: max_pages * 80],
            source="native_docx",
            confidence=0.9 if lines else 0.2,
        ),
    ]


def _extract_image_tesseract(data: bytes, max_pages: int) -> list[OcrPageResult]:
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        return [_error_page("ocr_unavailable")]

    try:
        img = Image.open(io.BytesIO(data))
        text = pytesseract.image_to_string(img, lang="rus+eng")
    except Exception:
        return [_error_page("ocr_failed")]
    lines_raw = [ln for ln in text.splitlines() if ln.strip()]
    y = 40.0
    lines: list[tuple[str, float, float, float, float]] = []
    for ln in lines_raw[: max_pages * 80]:
        lines.append((ln, 40.0, y, float(img.width - 80), 16.0))
        y += 18.0
    return [
        _lines_to_page(
            page_number=1,
            width=float(img.width),
            height=float(img.height),
            lines=lines,
            source="ocr",
            confidence=0.75 if lines else 0.2,
        ),
    ]


def extract_document_pages(
    data: bytes,
    *,
    detected_type: str,
    content_type: str | None = None,
    max_pages: int = MAX_PAGES_DEFAULT,
) -> OcrDocumentResult:
    """Extract pages from real file bytes. Never uses fixture ids."""
    dt = (detected_type or "").lower()
    ct = (content_type or "").lower()
    pages: list[OcrPageResult] = []

    if dt in {"pdf", "application/pdf"} or "pdf" in ct:
        try:
            import fitz  # noqa: F401

            pages = _extract_pdf_pymupdf(data, max_pages)
        except Exception:
            pages = _extract_pdf_pypdf(data, max_pages)
    elif dt in {"docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"} or "word" in ct:
        pages = _extract_docx(data, max_pages)
    elif dt in {"jpeg", "jpg", "png", "image/jpeg", "image/png"} or ct.startswith("image/"):
        pages = _extract_image_tesseract(data, max_pages)

    if not pages:
        pages = [
            _lines_to_page(
                page_number=1,
                width=595.0,
                height=842.0,
                lines=[],
                source="unknown",
                confidence=0.0,
            ),
        ]
        pages[0] = OcrPageResult(
            page_number=1,
            text="",
            confidence=0.0,
            width=595.0,
            height=842.0,
            rotation=0,
            language="ru",
            words=(),
            lines=(),
            blocks=(),
            source="unknown",
            error_code="unsupported_or_empty",
        )

    return OcrDocumentResult(pages=pages, provider="local_extract", model_version="local-extract-2")


def page_count_limit_exceeded(data: bytes, detected_type: str, max_pages: int) -> bool:
    if detected_type != "pdf":
        return False
    try:
        from pypdf import PdfReader

        return len(PdfReader(io.BytesIO(data)).pages) > max_pages
    except Exception:
        return False


def estimate_page_count(data: bytes, detected_type: str) -> int | None:
    if detected_type != "pdf":
        return 1 if data else 0
    try:
        from pypdf import PdfReader

        return len(PdfReader(io.BytesIO(data)).pages)
    except Exception:
        return None
